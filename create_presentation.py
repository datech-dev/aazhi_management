import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_aazhi_presentation(output_path):
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]  # Blank layout
    
    # Color Palette Definitions
    BG_DARK = RGBColor(17, 24, 39)       # Deep Slate / Midnight Navy (#111827)
    CARD_BG = RGBColor(31, 41, 55)       # Card background (#1F2937)
    GOLD_ACCENT = RGBColor(212, 175, 55) # Luxury Gold (#D4AF37)
    GOLD_LIGHT = RGBColor(243, 217, 137) # Light Warm Gold (#F3D989)
    TEXT_WHITE = RGBColor(255, 255, 255) # Pure White
    TEXT_MUTED = RGBColor(209, 213, 219)# Light Silver Grey (#D1D5DB)
    ACCENT_EMERALD = RGBColor(16, 185, 129) # Emerald Green (#10B981)
    ACCENT_ROSE = RGBColor(244, 63, 94)     # Rose Pink (#F43F5E)
    
    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color
        
    def add_header(slide, title_text, category_text="AAZHI DESIGNER STUDIO — DEMO PRESENTATION"):
        # Category / Kicker text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GOLD_ACCENT
        p_cat.font.name = 'Georgia'
        
        # Main Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = 'Georgia'
        
        # Decorative divider line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.03)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD_ACCENT
        line.line.fill.background()

    def add_card(slide, left, top, width, height, title, items, highlight_tag=None, card_bg=CARD_BG, border_color=GOLD_ACCENT):
        # Card container box
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = card_bg
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        # Text Frame inside card
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        if highlight_tag:
            p_tag = tf.paragraphs[0]
            p_tag.text = highlight_tag.upper()
            p_tag.font.size = Pt(10)
            p_tag.font.bold = True
            p_tag.font.color.rgb = GOLD_ACCENT
            p_tag.space_after = Pt(4)
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]
            
        p_title.text = title
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = GOLD_LIGHT
        p_title.font.name = 'Georgia'
        p_title.space_after = Pt(12)
        
        for item in items:
            p_item = tf.add_paragraph()
            p_item.space_after = Pt(8)
            p_item.font.size = Pt(13)
            p_item.font.name = 'Calibri'
            
            if isinstance(item, tuple):
                # (bold_label, text_desc)
                run1 = p_item.add_run()
                run1.text = "• " + item[0] + ": "
                run1.font.bold = True
                run1.font.color.rgb = TEXT_WHITE
                
                run2 = p_item.add_run()
                run2.text = item[1]
                run2.font.color.rgb = TEXT_MUTED
            else:
                run = p_item.add_run()
                run.text = "• " + item
                run.font.color.rgb = TEXT_MUTED

    def add_footer(slide, current_slide, total_slides=13):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Aazhi Designer Studio — Confidential Customer Demo Presentation | Slide {current_slide} of {total_slides}"
        p.font.size = Pt(10)
        p.font.color.rgb = RGBColor(156, 163, 175)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Cover)
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, BG_DARK)
    
    # Outer decorative frame
    frame = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.333), Inches(6.5))
    frame.fill.background()
    frame.line.color.rgb = GOLD_ACCENT
    frame.line.width = Pt(2)
    
    # Large Title
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "🪡 AAZHI DESIGNER STUDIO"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = GOLD_LIGHT
    p1.font.name = 'Georgia'
    p1.space_after = Pt(10)
    
    p2 = tf1.add_paragraph()
    p2.text = "Boutique Business Management & Automation System"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.font.name = 'Georgia'
    p2.space_after = Pt(20)
    
    p3 = tf1.add_paragraph()
    p3.text = "Transforming Bespoke Fashion, Social Commerce & Tailoring Operations from Inquiry to Delivery"
    p3.font.size = Pt(16)
    p3.font.color.rgb = TEXT_MUTED
    p3.font.name = 'Calibri'
    
    # Subtitle details box
    det_box = slide1.shapes.add_textbox(Inches(1.2), Inches(4.8), Inches(10.9), Inches(1.5))
    tf_det = det_box.text_frame
    
    p_d1 = tf_det.paragraphs[0]
    p_d1.text = "✨ Product Demonstration & Solution Overview for Prospective Clients"
    p_d1.font.size = Pt(14)
    p_d1.font.bold = True
    p_d1.font.color.rgb = GOLD_ACCENT
    p_d1.space_after = Pt(6)
    
    p_d2 = tf_det.add_paragraph()
    p_d2.text = "Unified Omnichannel Inbox • Bespoke Garment Measurement Engine • Production Kanban • Financial Governance"
    p_d2.font.size = Pt(13)
    p_d2.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 2: Industry Challenges & Pain Points
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_DARK)
    add_header(slide2, "Industry Problem: Operational Bottlenecks in Luxury Boutiques")
    
    add_card(slide2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.9), 
             "📱 Social Media Chaos", 
             [("Fragmented DMs", "Inquiries scattered across WhatsApp & Instagram DM."),
              ("Slow Lead Response", "Potential high-value customers lost due to delayed replies."),
              ("No Follow-up Tracking", "Staff forget unconfirmed leads without a central CRM.")],
             highlight_tag="Customer Acquisition Loss")
             
    add_card(slide2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.9), 
             "✂️ Measurement & Fitting Errors", 
             [("Paper Card Losses", "Measurements recorded on physical paper books get misplaced."),
              ("Garment Fitting Errors", "Lack of versioning leads to stitching based on outdated measurements."),
              ("Expensive Rework", "Altering improperly tailored designer outfits eats profit margins.")],
             highlight_tag="Quality & Production Failure")

    add_card(slide2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.9), 
             "💳 Financial & Status Blindspots", 
             [("Uncollected Balances", "Outfits delivered before collecting remaining 50% balance."),
              ("Opaque Workshop Status", "Clients constantly call asking: 'Is my blouse ready?'"),
              ("Inventory Leakage", "Fabric rolls, embellishments, and trims untracked.")],
             highlight_tag="Revenue & Overhead Leakage")
             
    add_footer(slide2, 2)

    # -------------------------------------------------------------
    # SLIDE 3: The Complete Solution Overview
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_DARK)
    add_header(slide3, "The Solution: Aazhi Designer Studio Management Platform")
    
    # Left column: Architecture Overview Card
    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "🏆 All-in-One Operating System",
             [("Built Specifically for Fashion", "Tailored for Indian boutique fashion, bridal wear, blouses, and custom tailoring."),
              ("End-to-End Lifecycle", "Manages Lead -> Enquiry -> Quote -> Order -> Tailoring -> Quality Check -> Payment -> Delivery."),
              ("Multi-Branch Scalability", "Supports single-boutique owners or multi-branch designer chains with isolated data."),
              ("Role-Based Workflow", "Customized interfaces for Studio Owner, Sales Staff, and Master Tailors.")],
             highlight_tag="Core Value Proposition")

    # Right column: 4 Key Pillar Badges
    pillars = [
        ("💬 Unified Social Inbox", "WhatsApp & Instagram DM integration with automated order updates & template messaging."),
        ("📐 Dynamic Measurement Engine", "Garment-specific measurement templates (Blouse, Kurti, Saree, Lehenga) with versioning."),
        ("✂️ Live Production Kanban Board", "Step-by-step garment stage tracking: Cutting -> Stitching -> Finishing -> Quality Check."),
        ("💰 Decimal-Safe Financial Accounting", "INR (₹) advance payment tracking, balance enforcement, GST tax receipts, and profit reports.")
    ]
    
    top_pos = 1.8
    for title, desc in pillars:
        p_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(top_pos), Inches(5.733), Inches(1.05))
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = CARD_BG
        p_card.line.color.rgb = GOLD_ACCENT
        p_card.line.width = Pt(1)
        
        tb = slide3.shapes.add_textbox(Inches(6.9), Inches(top_pos + 0.05), Inches(5.533), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = GOLD_LIGHT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        
        top_pos += 1.25

    add_footer(slide3, 3)

    # -------------------------------------------------------------
    # SLIDE 4: Module 1 — Unified Social Inbox
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_DARK)
    add_header(slide4, "Module 1: Unified Social Inbox & Omnichannel Communication")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "📲 Social Media Hub",
             [("Single Window Inbox", "Consolidates messages from WhatsApp Business API and Instagram Direct Messages into one live feed."),
              ("Instant Lead Conversion", "Turn an Instagram inquiry about a bridal blouse directly into a CRM Lead or Order with one click."),
              ("Staff Assignment", "Route conversations to specific sales executives with priority tags (Urgent, High, VIP)."),
              ("Media Attachments", "Receive and store customer inspiration photos, fabric references, and sketch uploads.")],
             highlight_tag="Omnichannel Messaging")

    add_card(slide4, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "⚡ Automated Customer Updates",
             [("Template Engine", "Pre-approved templates for Order Confirmations, Payment Receipts, and Delivery Notifications."),
              ("Variable Substitution", "Auto-populates {{customer_name}}, {{order_number}}, {{balance_due}}, and {{delivery_date}}."),
              ("Status Triggers", "Moving an order to 'Ready for Pickup' instantly dispatches a WhatsApp notification to the client."),
              ("Zero Manual Effort", "Keeps clients informed effortlessly without constant phone calls.")],
             highlight_tag="Customer Engagement & Automation")

    add_footer(slide4, 4)

    # -------------------------------------------------------------
    # SLIDE 5: Module 2 — Customer 360° CRM & Lead Pipeline
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_DARK)
    add_header(slide5, "Module 2: Customer 360° CRM & Lead Pipeline")

    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "👥 Customer 360° Profile",
             [("Complete History", "Unified view of contact details, WhatsApp/IG handles, address book, and order logs."),
              ("Lifetime Value (LTV) Tracking", "Automatically calculates total spend and total orders per client."),
              ("VIP & Preferred Tags", "Tag high-spending clients for priority queueing and special seasonal discounts."),
              ("Special Event Tracking", "Records Customer Birthdays and Anniversaries for personalized marketing outreach.")],
             highlight_tag="Client Intelligence")

    add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "📊 Lead Management Funnel",
             [("Stage Tracking", "Move leads through: New -> Contacted -> Interested -> Quote Sent -> Negotiation -> Confirmed."),
              ("Estimated Deal Value", "Track potential revenue pipeline in real-time in INR (₹)."),
              ("Follow-Up Reminders", "Set automated follow-up alerts so sales staff never miss a interested bride or client."),
              ("Lead-to-Order Conversion", "1-click conversion copies customer metadata, measurements, and items directly into a formal order.")],
             highlight_tag="Sales Pipeline Control")

    add_footer(slide5, 5)

    # -------------------------------------------------------------
    # SLIDE 6: Module 3 — Bespoke Measurement Engine
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_DARK)
    add_header(slide6, "Module 3: Bespoke Measurement Engine")

    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "📐 Dynamic Measurement Templates",
             [("Garment-Specific Fields", "Custom templates for Blouses, Kurtis, Sarees, Lehengas, Anarkalis, and Gowns."),
              ("Comprehensive Metrics", "Bust, Underbust, Waist, Hips, Shoulder, Armhole, Bicep, Sleeve Length, Front/Back Neck Depth, Garment Length."),
              ("Inches & CM Support", "Seamlessly toggle between Inches and Centimeters per client preference."),
              ("Mandatory Validation", "Ensures tailors have all required measurements before cutting fabric.")],
             highlight_tag="Precision Fitting")

    add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "🔄 Versioning & Order Snapshotting",
             [("Historical Versioning", "Retains past measurement profiles as clients' fitting requirements change over time."),
              ("Order-Bound Profiles", "Locks exact measurement values to specific order IDs so alterations to profile don't corrupt past orders."),
              ("Design Customization Notes", "Capture neck line styles (Potli, Sweetheart, Boat neck), sleeve lining details, and padding specs."),
              ("Master Tailor Printouts", "Generate clean digitised measurement cards for master craftsmen in the workshop.")],
             highlight_tag="Zero-Error Tailoring")

    add_footer(slide6, 6)

    # -------------------------------------------------------------
    # SLIDE 7: Module 4 — Tailoring & Production Pipeline
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_DARK)
    add_header(slide7, "Module 4: Tailoring & Production Kanban Pipeline")

    # 5 Process Stages visually depicted
    stages = [
        ("1. Cutting ✂️", "Fabric inspection, pattern drafting & precision cutting by Master Cutter."),
        ("2. Stitching 🪡", "Garment assembly, embroidery placement & seam stitching by assigned Tailor."),
        ("3. Finishing 🧺", "Hemming, hook/potli attachment, ironing & steam pressing."),
        ("4. Quality Check ✔️", "Verification against measurement profile, neck depth, and design specs."),
        ("5. Ready & Delivery 🚚", "Final garment packaging, balance collection, and dispatch/pickup.")
    ]
    
    col_width = Inches(2.2)
    spacing = Inches(0.18)
    start_left = Inches(0.8)
    
    for i, (st_title, st_desc) in enumerate(stages):
        l_pos = start_left + i * (col_width + spacing)
        card_st = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_pos, Inches(1.8), col_width, Inches(4.9))
        card_st.fill.solid()
        card_st.fill.fore_color.rgb = CARD_BG
        card_st.line.color.rgb = GOLD_ACCENT if i == 3 else RGBColor(75, 85, 99)
        card_st.line.width = Pt(1.5 if i == 3 else 1)
        
        tb = slide7.shapes.add_textbox(l_pos + Inches(0.1), Inches(1.9), col_width - Inches(0.2), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_st = tf.paragraphs[0]
        p_st.text = st_title
        p_st.font.size = Pt(14)
        p_st.font.bold = True
        p_st.font.color.rgb = GOLD_LIGHT
        p_st.space_after = Pt(10)
        
        p_sd = tf.add_paragraph()
        p_sd.text = st_desc
        p_sd.font.size = Pt(11)
        p_sd.font.color.rgb = TEXT_MUTED
        p_sd.space_after = Pt(14)
        
        p_feat = tf.add_paragraph()
        if i == 0:
            p_feat.text = "• Assigns Master Cutter\n• Pattern sheet print"
        elif i == 1:
            p_feat.text = "• Tailor task queue\n• Work distribution"
        elif i == 2:
            p_feat.text = "• Trims & lining check\n• Steam press status"
        elif i == 3:
            p_feat.text = "• Metric verification\n• Pass / Alteration routing"
        else:
            p_feat.text = "• Auto WhatsApp alert\n• Balance collection check"
        p_feat.font.size = Pt(10)
        p_feat.font.color.rgb = TEXT_WHITE

    add_footer(slide7, 7)

    # -------------------------------------------------------------
    # SLIDE 8: Module 5 — Quoting & Order Lifecycle
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, BG_DARK)
    add_header(slide8, "Module 5: Quotations & Order Lifecycle Management")

    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "📝 Professional Quotations",
             [("Digital Quotation Builder", "Create itemized price quotes with fabric costs, tailoring charges, and embroidery add-ons."),
              ("Validity Period Control", "Set quote expiration dates (e.g. 7 days) to protect against fluctuating silk/fabric prices."),
              ("PDF / WhatsApp Dispatch", "Send elegant PDF quotations directly to clients over WhatsApp or Email."),
              ("Instant Order Conversion", "Client accepts quote -> Convert to live Order instantly without re-keying items.")],
             highlight_tag="Sales Enablement")

    add_card(slide8, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "📦 Order Lifecycle Governance",
             [("Unique Sequence Numbering", "Automatic branch-specific order codes (e.g. AZ-2026-0001)."),
              ("Promised Delivery Dates", "Tracks target delivery deadlines with urgency badges (Low, Medium, High, Urgent)."),
              ("Line Item Customization", "Detailed notes per garment item (e.g. 'Add potli buttons on back, padded blouse')."),
              ("Audit History", "Complete timestamped audit log of status changes and staff edits.")],
             highlight_tag="Operational Precision")

    add_footer(slide8, 8)

    # -------------------------------------------------------------
    # SLIDE 9: Module 6 — Financial Management & Revenue Governance
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, BG_DARK)
    add_header(slide9, "Module 6: Financial Management & Revenue Governance")

    add_card(slide9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "💰 Decimal-Safe Financial Math",
             [("INR (₹) Precision Math", "Built with high-precision decimal math engine for accurate tax, discounts, and advance payments."),
              ("Advance Enforcer", "Require configurable advance percentage (e.g. 50%) before order transitions to cutting."),
              ("Balance Due Controls", "Alerts staff if an order is marked for delivery with an unpaid balance."),
              ("Payment Methods", "Records Cash, UPI (GPay/PhonePe), Card, and Bank Transfer receipts.")],
             highlight_tag="Zero Revenue Leakage")

    add_card(slide9, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "📊 Financial Reports & Analytics",
             [("Revenue Dashboards", "Daily, weekly, and monthly gross revenue, collected payments, and pending balances."),
              ("Payment Status Breakdown", "Filter orders by Unpaid, Advance Paid, Partially Paid, and Fully Paid."),
              ("GST Tax Receipts", "Generate compliant invoices with customizable GST/Tax percentages."),
              ("Audit-Proof Ledger", "Immutable payment record logs linked to staff user accounts.")],
             highlight_tag="Business Intelligence")

    add_footer(slide9, 9)

    # -------------------------------------------------------------
    # SLIDE 10: Module 7 — Catalog, Inventory & Materials
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, BG_DARK)
    add_header(slide10, "Module 7: Product Catalog & Material Inventory")

    add_card(slide10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "👗 Product & Designer Catalog",
             [("Outfit Categories", "Organize collections into Sarees, Lehengas, Designer Blouses, Indo-Western, and Fabrics."),
              ("SKU & Variant Management", "Track stock by Size (S, M, L, XL), Color, and Fabric composition."),
              ("Customizable Flags", "Mark items as 'Customizable Garment' or 'Ready-to-Wear'."),
              ("High-Res Gallery", "Store image catalogs for staff to share with clients during sales consultations.")],
             highlight_tag="Catalog Management")

    add_card(slide10, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "🧵 Material & Trim Inventory",
             [("Raw Material Tracking", "Monitor yards of raw silk, organza, net, linings, and laces in stock."),
              ("Low Stock Alerts", "Automated alerts when fabric rolls fall below reorder thresholds."),
              ("Cost Price vs Sale Price", "Track material margins and profit profitability per garment type."),
              ("Multi-Branch Stock Isolation", "Manage inventory across main store and workshop units seamlessly.")],
             highlight_tag="Stock & Cost Control")

    add_footer(slide10, 10)

    # -------------------------------------------------------------
    # SLIDE 11: Enterprise Security, RBAC & Multi-Branch Support
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, BG_DARK)
    add_header(slide11, "Enterprise Governance: Multi-Branch & Granular RBAC")

    add_card(slide11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "🏢 Multi-Branch Isolation",
             [("Branch Architecture", "Isolated data access per store location (e.g. Master Store vs Branch 2)."),
              ("Branch Sequence Codes", "Unique order & lead numbering per branch (`AZ-MAS-001`, `AZ-B2-001`)."),
              ("Central Owner Dashboard", "Studio Owner accesses cross-branch aggregated sales reports in real time."),
              ("Staff Branch Binding", "Staff members restricted to their assigned branch operations.")],
             highlight_tag="Multi-Location Scale")

    add_card(slide11, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "🔐 Granular Role-Based Access",
             [("Studio Owner 👑", "Full system access, financial reports, pricing overrides, and staff management."),
              ("Sales Executive 🛍️", "Lead creation, customer management, quote generation, and order creation."),
              ("Master Tailor ✂️", "Restricted workshop view: Production Kanban, measurement cards, and job status."),
              ("Bank-Grade Security", "Bcrypt password hashing, JWT sessions, and server-enforced permission guards.")],
             highlight_tag="Access Control & Security")

    add_footer(slide11, 11)

    # -------------------------------------------------------------
    # SLIDE 12: Business Impact & Expected ROI
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, BG_DARK)
    add_header(slide12, "Why Choose Aazhi: Quantifiable Business Impact & ROI")

    metrics = [
        ("⚡ 40% Faster Lead Response", "Convert Instagram & WhatsApp inquiries into paying customers rapidly before competitors reply."),
        ("🎯 95% Drop in Fitting Errors", "Digitized garment measurement profiles eliminate re-stitching costs and client disappointment."),
        ("💰 100% Balance Recovery", "Strict financial balance enforcement prevents orders from leaving workshop un-paid."),
        ("📈 30% Boost in Tailor Efficiency", "Clear Kanban stage assignment prevents workshop bottlenecks and idle tailor time.")
    ]

    top_pos = 1.8
    for title, desc in metrics:
        m_card = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.733), Inches(1.05))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = CARD_BG
        m_card.line.color.rgb = GOLD_ACCENT
        m_card.line.width = Pt(1.5)
        
        tb = slide12.shapes.add_textbox(Inches(1.0), Inches(top_pos + 0.1), Inches(11.333), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = GOLD_LIGHT
        p1.space_after = Pt(2)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED
        
        top_pos += 1.25

    add_footer(slide12, 12)

    # -------------------------------------------------------------
    # SLIDE 13: Live Demo Credentials & Next Steps
    # -------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, BG_DARK)
    add_header(slide13, "Live Demonstration Walkthrough & Next Steps")

    add_card(slide13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9),
             "🔑 Demo Account Credentials",
             [("Studio Owner Role 👑", "Email: owner@aazhi.studio | Pass: Aazhi@2026!\nAccess: Full dashboard, financial reports & settings."),
              ("Sales Executive Role 🛍️", "Email: sales@aazhi.studio | Pass: Aazhi@2026!\nAccess: Social Inbox, CRM, Quotes & Orders."),
              ("Master Tailor Role ✂️", "Email: tailor@aazhi.studio | Pass: Aazhi@2026!\nAccess: Production Kanban & Measurement Sheets.")],
             highlight_tag="Interactive Walkthrough")

    add_card(slide13, Inches(6.8), Inches(1.8), Inches(5.733), Inches(4.9),
             "🚀 Onboarding & Next Steps",
             [("Step 1: Live Demo Walkthrough", "Experience live WhatsApp inbound message flow and tailoring job assignment."),
              ("Step 2: Custom Measurement Setup", "Configure your boutique's specific garment templates and size charts."),
              ("Step 3: Staff Role Provisioning", "Create accounts for your studio managers, sales team, and master tailors."),
              ("Step 4: Go-Live & Cloud Hosting", "Deploy on high-performance secure cloud infrastructure with automated daily backups.")],
             highlight_tag="Getting Started")

    add_footer(slide13, 13)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "Aazhi_Designer_Studio_Customer_Demo.pptx")
    create_aazhi_presentation(out_path)
